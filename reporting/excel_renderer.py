import logging
import json
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger("reconciliation_agent.reporting.excel_renderer")

def apply_auto_width(ws):
    """Adjust worksheet column widths to fit contents."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Check cell value length
            val = str(cell.value or '')
            # If it looks like a long JSON block, cap the width calculation
            if len(val) > 40:
                val = val[:40]
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

def render_excel_report(session_state: dict, output_dir: Path) -> Path:
    """Generates a styled, multi-sheet Excel reconciliation report.
    
    Returns:
        The path of the generated Excel file.
    """
    logger.info("Starting Excel report generation")
    
    results = session_state['reconciliation_results']
    primary_keys = session_state['primary_key_cols']
    source_filename = session_state['source_filename']
    target_filename = session_state['target_filename']
    
    wb = Workbook()
    
    # ── Styling Config ──
    font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    font_title = Font(name="Arial", size=14, bold=True, color="312E81")
    font_bold = Font(name="Arial", size=10, bold=True)
    font_regular = Font(name="Arial", size=10)
    
    fill_indigo = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    fill_gray_light = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    
    # Severity fills
    fill_critical = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # light red
    fill_warning = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")  # light yellow
    fill_info = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")     # light blue
    
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    # 1. Summary Sheet
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_sum.views.sheetView[0].showGridLines = True
    
    ws_sum["A1"] = "Data Reconciliation Report"
    ws_sum["A1"].font = font_title
    
    metadata = [
        ("Source File", source_filename),
        ("Target File", target_filename),
        ("Reconciliation Keys", ", ".join(primary_keys)),
        ("Execution Date", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    
    ws_sum.append([]) # spacing
    for label, val in metadata:
        ws_sum.append([label, val])
        row_idx = ws_sum.max_row
        ws_sum[f"A{row_idx}"].font = font_bold
        ws_sum[f"B{row_idx}"].font = font_regular
        
    ws_sum.append([]) # spacing
    
    # Table Header for metrics
    ws_sum.append(["Reconciliation Metrics", "Count"])
    header_row = ws_sum.max_row
    ws_sum[f"A{header_row}"].font = font_header
    ws_sum[f"A{header_row}"].fill = fill_indigo
    ws_sum[f"B{header_row}"].font = font_header
    ws_sum[f"B{header_row}"].fill = fill_indigo
    
    metrics = [
        ("Total Rows (Source)", results['summary']['total_source_rows']),
        ("Total Rows (Target)", results['summary']['total_target_rows']),
        ("Clean Source Rows", results['summary']['clean_source_rows']),
        ("Clean Target Rows", results['summary']['clean_target_rows']),
        ("Matched Rows", results['summary']['matched_rows']),
        ("Mismatched Rows", results['summary']['mismatched_rows']),
        ("Missing in Target", results['summary']['missing_in_target']),
        ("Missing in Source", results['summary']['missing_in_source']),
        ("Duplicate Rows (Source)", results['summary']['duplicate_source']),
        ("Duplicate Rows (Target)", results['summary']['duplicate_target'])
    ]
    
    for metric, count in metrics:
        ws_sum.append([metric, count])
        row_idx = ws_sum.max_row
        ws_sum[f"A{row_idx}"].font = font_regular
        ws_sum[f"B{row_idx}"].font = font_bold
        ws_sum[f"B{row_idx}"].alignment = Alignment(horizontal="right")
        ws_sum[f"A{row_idx}"].border = thin_border
        ws_sum[f"B{row_idx}"].border = thin_border

    # 2. Mismatches Sheet
    ws_mis = wb.create_sheet(title="Mismatches")
    ws_mis.views.sheetView[0].showGridLines = True
    
    headers_mis = ["Row Key", "Column Name", "Source Value", "Target Value", "Severity"]
    ws_mis.append(headers_mis)
    for col_idx in range(1, 6):
        cell = ws_mis.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_indigo
        
    for item in results["mismatches"]:
        ws_mis.append([
            item["row_key"],
            item["column"],
            item["source_value"],
            item["target_value"],
            item["severity"]
        ])
        
        row_idx = ws_mis.max_row
        sev = item["severity"]
        
        # Determine fill based on severity
        fill = fill_info
        if sev == "Critical":
            fill = fill_critical
        elif sev == "Warning":
            fill = fill_warning
            
        # Apply fonts, fills, borders
        for col_idx in range(1, 6):
            cell = ws_mis.cell(row=row_idx, column=col_idx)
            cell.font = font_regular
            cell.fill = fill
            cell.border = thin_border
            if col_idx in (3, 4):
                cell.font = Font(name="Courier New", size=10)

    # 3. Missing in Target Sheet
    ws_mt = wb.create_sheet(title="Missing In Target")
    ws_mt.views.sheetView[0].showGridLines = True
    ws_mt.append(["Row Key", "Record Details (JSON)"])
    for col_idx in range(1, 3):
        cell = ws_mt.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_indigo
        
    for item in results["missing_in_target"]:
        ws_mt.append([item["row_key"], json.dumps(item["row_data"])])
        row_idx = ws_mt.max_row
        ws_mt.cell(row=row_idx, column=1).font = font_bold
        ws_mt.cell(row=row_idx, column=1).border = thin_border
        ws_mt.cell(row=row_idx, column=2).font = Font(name="Courier New", size=9)
        ws_mt.cell(row=row_idx, column=2).border = thin_border

    # 4. Missing in Source Sheet
    ws_ms = wb.create_sheet(title="Missing In Source")
    ws_ms.views.sheetView[0].showGridLines = True
    ws_ms.append(["Row Key", "Record Details (JSON)"])
    for col_idx in range(1, 3):
        cell = ws_ms.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_indigo
        
    for item in results["missing_in_source"]:
        ws_ms.append([item["row_key"], json.dumps(item["row_data"])])
        row_idx = ws_ms.max_row
        ws_ms.cell(row=row_idx, column=1).font = font_bold
        ws_ms.cell(row=row_idx, column=1).border = thin_border
        ws_ms.cell(row=row_idx, column=2).font = Font(name="Courier New", size=9)
        ws_ms.cell(row=row_idx, column=2).border = thin_border

    # 5. Duplicates Source Sheet (only if duplicates exist)
    if len(results["duplicates_source"]) > 0:
        ws_ds = wb.create_sheet(title="Duplicates Source")
        ws_ds.views.sheetView[0].showGridLines = True
        ws_ds.append(["Row Key", "Full Record JSON"])
        for col_idx in range(1, 3):
            cell = ws_ds.cell(row=1, column=col_idx)
            cell.font = font_header
            cell.fill = fill_indigo
            
        for item in results["duplicates_source"]:
            row_key = item.get("_row_key", "")
            clean_item = {k: v for k, v in item.items() if k != "_row_key"}
            ws_ds.append([row_key, json.dumps(clean_item)])
            row_idx = ws_ds.max_row
            ws_ds.cell(row=row_idx, column=1).font = font_bold
            ws_ds.cell(row=row_idx, column=1).border = thin_border
            ws_ds.cell(row=row_idx, column=2).font = Font(name="Courier New", size=9)
            ws_ds.cell(row=row_idx, column=2).border = thin_border

    # 6. Duplicates Target Sheet (only if duplicates exist)
    if len(results["duplicates_target"]) > 0:
        ws_dt = wb.create_sheet(title="Duplicates Target")
        ws_dt.views.sheetView[0].showGridLines = True
        ws_dt.append(["Row Key", "Full Record JSON"])
        for col_idx in range(1, 3):
            cell = ws_dt.cell(row=1, column=col_idx)
            cell.font = font_header
            cell.fill = fill_indigo
            
        for item in results["duplicates_target"]:
            row_key = item.get("_row_key", "")
            clean_item = {k: v for k, v in item.items() if k != "_row_key"}
            ws_dt.append([row_key, json.dumps(clean_item)])
            row_idx = ws_dt.max_row
            ws_dt.cell(row=row_idx, column=1).font = font_bold
            ws_dt.cell(row=row_idx, column=1).border = thin_border
            ws_dt.cell(row=row_idx, column=2).font = Font(name="Courier New", size=9)
            ws_dt.cell(row=row_idx, column=2).border = thin_border

    # Adjust column widths for all sheets
    for sheet in wb.worksheets:
        apply_auto_width(sheet)
        
    # Generate timestamped filename
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    output_file = output_dir / f"recon_{timestamp}.xlsx"
    wb.save(output_file)
    
    logger.info("Successfully generated Excel report at %s", output_file)
    return output_file
