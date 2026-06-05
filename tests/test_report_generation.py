import pytest
from pathlib import Path
from tools import SESSION_STATE
from tools.run_reconciliation import run_reconciliation
from tools.generate_report import generate_report
from utils.data_normalizer import align_columns

def test_generate_report_no_results():
    res = generate_report.func("html")
    assert "Error" in res

def test_generate_report_html_and_excel(tmp_path, mock_source_df, mock_target_df):
    # Setup state
    comp_src, comp_tgt, meta = align_columns(mock_source_df, mock_target_df)
    SESSION_STATE.comp_source = comp_src
    SESSION_STATE.comp_target = comp_tgt
    SESSION_STATE.source_df = mock_source_df
    SESSION_STATE.target_df = mock_target_df
    SESSION_STATE.align_meta = meta
    SESSION_STATE.source_filename = "source.csv"
    SESSION_STATE.target_filename = "target.json"
    SESSION_STATE.source_fullpath = "/dummy/source.csv"
    SESSION_STATE.target_fullpath = "/dummy/target.json"
    
    # Run reconciliation comparison
    run_reconciliation.func("id")
    
    # Override REPORTS_DIR for the test run so we output to tmp_path
    import tools.generate_report
    # Save the original directory and temporarily patch it
    orig_dir = tools.generate_report.REPORTS_DIR
    tools.generate_report.REPORTS_DIR = tmp_path
    
    try:
        # Test HTML report generation
        res_html = generate_report.func("html")
        assert "Success!" in res_html
        
        # Verify HTML file is created
        html_files = list(tmp_path.glob("*.html"))
        assert len(html_files) == 1
        html_content = html_files[0].read_text(encoding="utf-8")
        assert "Data Reconciliation Report" in html_content
        assert "ReconAgent" in html_content
        assert "window.FULL_REPORT_DATA" in html_content
        
        # Test Excel report generation
        res_xlsx = generate_report.func("excel")
        assert "Success!" in res_xlsx
        
        # Verify Excel file is created
        xlsx_files = list(tmp_path.glob("*.xlsx"))
        assert len(xlsx_files) == 1
        
        # Open with openpyxl to check sheets
        from openpyxl import load_workbook
        wb = load_workbook(xlsx_files[0])
        sheet_names = wb.sheetnames
        assert "Summary" in sheet_names
        assert "Mismatches" in sheet_names
        assert "Missing In Target" in sheet_names
        assert "Missing In Source" in sheet_names
        
    finally:
        # Restore patched REPORTS_DIR
        tools.generate_report.REPORTS_DIR = orig_dir
